"""Excel export of a reconciliation, laid out like a signed file copy.

The statement sheet uses real Excel formulas for every subtotal, adjusted
balance and the difference, so anyone checking it in Excel can see how each
figure is built up and recalculates if they change a number.

Text from the uploaded files (descriptions, payees, reasons) is always written
as plain text, never as a formula. A bank description such as
=HYPERLINK("http://evil",...) could otherwise run when the file is opened
(this is called "CSV/formula injection").
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .money import ZERO
from .service import from_cents
from .timefmt import local_time

MONEY = '#,##0.00;(#,##0.00);"-"'
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
MUTED = Font(color="6B7078")
HEADER_FILL = PatternFill("solid", fgColor="17332D")
HEADER_FONT = Font(bold=True, color="FFFFFF")
FLAG_FILL = PatternFill("solid", fgColor="FFFAEB")
TOP_RULE = Border(top=Side(style="thin"))
TOTAL_RULE = Border(top=Side(style="thin"), bottom=Side(style="double"))
DANGEROUS_START = ("=", "+", "-", "@", "\t", "\r")


def put_text(ws, row, col, value, font=None):
    """Write text that can never be treated as a formula."""
    cell = ws.cell(row=row, column=col)
    cell.value = "" if value is None else str(value)
    if cell.value.startswith(DANGEROUS_START):
        cell.data_type = "s"           # stored as text, shown as typed
        cell.quotePrefix = True        # Excel won't reinterpret it when the cell is edited
    if font:
        cell.font = font
    return cell


def put_money(ws, row, col, value, font=None, border=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.number_format = MONEY
    if font:
        cell.font = font
    if border:
        cell.border = border
    return cell


def build_workbook(svc, run_id):
    run = svc.run(run_id)
    s = svc.statement(run_id)
    wb = Workbook()
    _statement_sheet(wb.active, run, s)
    _matches_sheet(wb.create_sheet("Matched items"), svc, run_id)
    _audit_sheet(wb.create_sheet("Audit trail"), svc.audit_trail(run_id))
    return wb


def workbook_bytes(svc, run_id):
    buffer = io.BytesIO()
    build_workbook(svc, run_id).save(buffer)
    return buffer.getvalue()


def file_name(run):
    return f"reconciliation-{run['last_digits']}-{run['period_end']}.xlsx"


# ---- Sheet 1: the statement ----------------------------------------------------

def _statement_sheet(ws, run, s):
    ws.title = "Reconciliation"
    for col, width in zip("ABCD", (13, 58, 16, 18)):
        ws.column_dimensions[col].width = width

    put_text(ws, 1, 1, "Bank Reconciliation Statement", TITLE)
    put_text(ws, 2, 1, f"{run['account_name']} · {run['bank_name']} account ending {run['last_digits']}")
    put_text(ws, 3, 1, f"As at {s.period_end.strftime('%d %B %Y')}", BOLD)
    put_text(ws, 4, 1, f"Status: {run['status'].capitalize()}", MUTED)
    row = 6

    # Bank side
    put_text(ws, row, 1, "Balance per bank statement")
    bank_start = put_money(ws, row, 4, s.bank_balance).coordinate
    row += 1
    parts = [bank_start]
    for title, items, flagged in (
            ("Add: deposits in transit", s.deposits_in_transit, False),
            ("Less: outstanding payments", s.outstanding_payments, False),
            ("Under investigation (books)", s.investigate_ledger, True)):
        row, cell = _item_section(ws, row, title, [(x.date, _ledger_text(x), x.amount) for x in items], flagged)
        if cell:
            parts.append(cell)
    row, cell = _item_section(ws, row, "Bank errors to be corrected by the bank",
                              [(None, label, amount) for label, amount in s.bank_corrections], False)
    if cell:
        parts.append(cell)
    put_text(ws, row, 1, "Adjusted bank balance", BOLD)
    adjusted_bank = put_money(ws, row, 4, "=" + "+".join(parts), BOLD, TOTAL_RULE).coordinate
    row += 2

    # Book side
    put_text(ws, row, 1, "Balance per books")
    book_start = put_money(ws, row, 4, s.book_balance).coordinate
    row += 1
    parts = [book_start]
    for title, items, flagged in (
            ("Add: bank credits not yet recorded", s.bank_credits_not_in_books, False),
            ("Less: bank debits not yet recorded", s.bank_debits_not_in_books, False),
            ("Under investigation (bank)", s.investigate_bank, True)):
        row, cell = _item_section(ws, row, title, [(x.date, _bank_text(x), x.amount) for x in items], flagged)
        if cell:
            parts.append(cell)
    row, cell = _item_section(ws, row, "Corrections to post in the books",
                              [(None, label, amount) for label, amount in s.book_corrections], False)
    if cell:
        parts.append(cell)
    put_text(ws, row, 1, "Adjusted book balance", BOLD)
    adjusted_book = put_money(ws, row, 4, "=" + "+".join(parts), BOLD, TOTAL_RULE).coordinate
    row += 2

    put_text(ws, row, 1, "Difference", BOLD)
    put_money(ws, row, 4, f"={adjusted_bank}-{adjusted_book}", BOLD).number_format = "#,##0.00;(#,##0.00);0.00"
    put_text(ws, row, 2, "Balances" if s.balances else "DOES NOT BALANCE", BOLD)
    row += 1
    if s.issues:
        row += 1
        put_text(ws, row, 1, f"Open issues ({len(s.issues)})", BOLD)
        for issue in s.issues:
            row += 1
            put_text(ws, row, 2, f"{issue.kind.replace('_', ' ')}: {issue.detail}")
        row += 1

    # Sign-off block
    row += 2
    put_text(ws, row, 1, "Prepared by", BOLD)
    put_text(ws, row, 2, f"{run['preparer_name']}   {local_time(run['submitted_at'])}")
    row += 1
    put_text(ws, row, 1, "Approved by", BOLD)
    put_text(ws, row, 2, f"{run['approver_name']}   {local_time(run['approved_at'])}"
             if run["approver_name"] else "Not yet approved")
    ws.freeze_panes = "A6"
    ws.print_options.gridLines = False
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0


def _item_section(ws, row, title, items, flagged):
    """Heading, one row per item, then a subtotal. Returns (next row, subtotal cell)."""
    if not items:
        return row, None
    put_text(ws, row, 1, title)
    first = row + 1
    for when, text, amount in items:
        row += 1
        date_cell = ws.cell(row=row, column=1, value=when)
        date_cell.number_format = "dd mmm yyyy"
        date_cell.alignment = Alignment(horizontal="left")
        put_text(ws, row, 2, text).alignment = Alignment(wrap_text=True, vertical="top")
        put_money(ws, row, 3, amount)
        if flagged:
            for col in range(1, 4):
                ws.cell(row=row, column=col).fill = FLAG_FILL
    last = row
    subtotal = put_money(ws, last, 4, f"=SUM(C{first}:C{last})")
    ws.cell(row=last, column=3).border = Border(bottom=Side(style="thin"))
    return row + 1, subtotal.coordinate


def _ledger_text(x):
    return " · ".join(p for p in (x.payee, x.voucher, x.reference) if p)


def _bank_text(x):
    return " · ".join(p for p in (x.description, x.reference) if p)


# ---- Sheet 2 and 3: supporting detail ---------------------------------------------

def _header(ws, columns, widths):
    for col, (name, width) in enumerate(zip(columns, widths), start=1):
        cell = put_text(ws, 1, col, name, HEADER_FONT)
        cell.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"


def _matches_sheet(ws, svc, run_id):
    _header(ws, ["Rule", "Bank items", "Bank amount", "Book items", "Book amount", "Difference", "Note", "Matched by"],
            [14, 40, 14, 40, 14, 12, 50, 18])
    bank = {r["line_no"]: r for r in svc.db.execute("SELECT * FROM bank_lines WHERE run_id = ?", (run_id,))}
    ledger = {r["line_no"]: r for r in svc.db.execute("SELECT * FROM ledger_entries WHERE run_id = ?", (run_id,))}
    for row, m in enumerate(svc.matches(run_id), start=2):
        bank_rows = [bank[n] for n in m["bank"]]
        ledger_rows = [ledger[n] for n in m["ledger"]]
        put_text(ws, row, 1, m["match"]["rule"].replace("_", " + "))
        put_text(ws, row, 2, "; ".join(r["description"] for r in bank_rows))
        put_money(ws, row, 3, sum((from_cents(r["amount"]) for r in bank_rows), ZERO))
        put_text(ws, row, 4, "; ".join(r["payee"] for r in ledger_rows))
        put_money(ws, row, 5, sum((from_cents(r["amount"]) for r in ledger_rows), ZERO))
        put_money(ws, row, 6, f"=C{row}-E{row}")
        put_text(ws, row, 7, m["match"]["note"])
        put_text(ws, row, 8, m["match"]["created_by_name"] or "Automatic")


def _audit_sheet(ws, entries):
    _header(ws, ["When (local time)", "Who", "Action", "Detail"], [18, 20, 18, 90])
    for row, a in enumerate(entries, start=2):
        put_text(ws, row, 1, local_time(a["at"]))
        put_text(ws, row, 2, a["user_name"])
        put_text(ws, row, 3, a["action"].replace("_", " "))
        put_text(ws, row, 4, a["detail"])
