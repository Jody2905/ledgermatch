"""Excel export: the formulas in the file must produce the same figures as the app."""

import io
import re
from datetime import datetime, timezone
from decimal import Decimal

from openpyxl import load_workbook

from ledgermatch import export
from ledgermatch.timefmt import local_time
from conftest import resolve_everything


def evaluate(ws, ref, depth=0):
    """Just enough of Excel to check our own formulas: numbers, =A1+B2-C3 and =SUM(C5:C9)."""
    assert depth < 50, "circular formula"
    value = ws[ref].value
    if not (isinstance(value, str) and value.startswith("=")):
        return Decimal(str(value or 0))
    expr = value[1:]

    def sum_range(m):
        col, start, end = m.group(1), int(m.group(2)), int(m.group(4))
        assert m.group(3) == col
        return "(" + "+".join(f"{col}{r}" for r in range(start, end + 1)) + ")"

    expr = re.sub(r"SUM\(([A-Z])(\d+):([A-Z])(\d+)\)", sum_range, expr)
    assert re.fullmatch(r"[A-Z0-9+\-()]+", expr), f"unexpected formula {value}"
    for cell in sorted(set(re.findall(r"[A-Z]+\d+", expr)), key=len, reverse=True):
        expr = re.sub(rf"\b{cell}\b", f"Decimal('{evaluate(ws, cell, depth + 1)}')", expr)
    return eval(expr, {"Decimal": Decimal})


def figure(ws, label):
    """The column-D figure on the row whose column-A label is `label`."""
    for row in ws.iter_rows(min_col=1, max_col=1):
        if row[0].value == label:
            return evaluate(ws, f"D{row[0].row}")
    raise AssertionError(f"{label!r} not found")


def load(svc, run_id):
    return load_workbook(io.BytesIO(export.workbook_bytes(svc, run_id)))


def test_formulas_reproduce_the_statement(svc, run_id):
    s = svc.statement(run_id)
    ws = load(svc, run_id)["Reconciliation"]
    assert figure(ws, "Adjusted bank balance") == s.adjusted_bank
    assert figure(ws, "Adjusted book balance") == s.adjusted_book
    assert figure(ws, "Difference") == 0
    assert any(c.value == "Open issues (2)" for c in ws["A"])


def test_resolved_reconciliation_shows_corrections_and_sign_off(svc, run_id, people):
    resolve_everything(svc, run_id, people["maker"])
    svc.submit(run_id, people["maker"])
    svc.approve(run_id, people["checker"])
    s = svc.statement(run_id)
    wb = load(svc, run_id)
    ws = wb["Reconciliation"]
    text = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert figure(ws, "Adjusted book balance") == s.adjusted_book == s.adjusted_bank
    assert "Corrections to post in the books" in text
    assert any("JV-2611" in t for t in text) and any("JV-2612" in t for t in text)
    assert "Status: Approved" in text
    assert any(t.startswith("Marcus James") for t in text)
    assert "Open issues" not in " ".join(text)
    assert wb.sheetnames == ["Reconciliation", "Matched items", "Audit trail"]


def test_matched_items_sheet(svc, run_id):
    ws = load(svc, run_id)["Matched items"]
    assert ws.max_row - 1 == len(svc.matches(run_id))
    for row in range(2, ws.max_row + 1):        # every automatic match has no difference
        assert evaluate(ws, f"F{row}") == 0


def test_audit_sheet_uses_local_time(svc, run_id):
    ws = load(svc, run_id)["Audit trail"]
    first = svc.audit_trail(run_id)[0]
    assert ws["A2"].value == local_time(first["at"])
    assert ws["B2"].value == "Kerry Baptiste"


def test_text_from_files_is_never_a_formula(svc, run_id, people):
    s = svc.statement(run_id)
    svc.mark_error(run_id, people["maker"], "ledger", s.investigate_ledger[0].line_no,
                   '=HYPERLINK("http://evil.example","click")')
    wb = load(svc, run_id)
    found = [c for ws in wb for row in ws.iter_rows() for c in row
             if isinstance(c.value, str) and "HYPERLINK" in c.value]
    assert found
    for cell in found:
        assert cell.data_type == "s"          # text, not formula


def test_file_name(svc, run_id):
    assert export.file_name(svc.run(run_id)) == "reconciliation-4821-2026-10-31.xlsx"


def test_local_time_converts_from_utc():
    stamp = "2026-10-11T04:59:00+00:00"
    expected = datetime(2026, 10, 11, 4, 59, tzinfo=timezone.utc).astimezone().strftime("%Y-%m-%d %H:%M")
    assert local_time(stamp) == expected
    assert local_time("2026-10-11T04:59:00") == expected     # no zone given: treated as UTC
    assert local_time(None) == ""
